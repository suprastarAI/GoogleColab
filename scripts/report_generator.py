"""
Generate laporan keuangan bulanan: markdown checklist + rencana pembayaran.
"""

from datetime import date
from typing import Optional
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

from config import CARDS, INDONESIAN_MONTHS, MONTHLY_BUDGET, DEFAULT_BUDGET
from calculator import CardPaymentPlan, get_active_installments_for_card


def _fmt(amount: int) -> str:
    return f"Rp {amount:,.0f}".replace(",", ".")


def _month_label(year_month: str) -> str:
    y, m = year_month.split("-")
    return f"{INDONESIAN_MONTHS[int(m)]} {y}"


def generate_checklist_md(
    plans: list[CardPaymentPlan],
    year_month: str,
    run_date: str,
) -> str:
    budget = MONTHLY_BUDGET.get(year_month, DEFAULT_BUDGET)
    label = _month_label(year_month)
    total_min = sum(p.minimum_payment for p in plans)
    total_rec = sum(p.recommended_payment for p in plans)
    total_bal = sum(p.current_balance for p in plans)

    lines = [
        f"# Checklist Pembayaran Kartu Kredit — {label}",
        f"**SUWANDHY PRAHARTO** | Diperbarui: {run_date} | Budget: **{_fmt(budget)}**",
        "",
        "---",
        "",
        "## Status Pembayaran",
        "",
        "| # | Kartu | Saldo | Min. Bayar | Rekomendasi | JT | Prioritas | Status |",
        "|---|---|---|---|---|---|---|---|",
    ]

    for i, p in enumerate(plans, 1):
        due = p.due_date or "—"
        est = " ⚠️est" if p.is_estimated else ""
        lines.append(
            f"| {i} | {p.bank} ...{p.last4} | {_fmt(p.current_balance)}{est} "
            f"| {_fmt(p.minimum_payment)} | {_fmt(p.recommended_payment)} "
            f"| {due} | {p.priority} | [ ] |"
        )

    lines += [
        f"| **TOTAL** | | **{_fmt(total_bal)}** | **{_fmt(total_min)}** | **{_fmt(total_rec)}** | | | |",
        "",
        "---",
        "",
        "## Kalkulasi Budget",
        "",
        "```",
        f"Budget {label:<30} {_fmt(budget):>18}",
        f"Total minimum payment         {_fmt(total_min):>18}",
        f"Total rekomendasi bayar       {_fmt(total_rec):>18}",
        f"Sisa budget (setelah rec.)    {_fmt(budget - total_rec):>18}",
        "```",
        "",
        "> ⚠️ Kartu bertanda **est** menggunakan estimasi — tagihan aktual belum masuk.",
        "",
        "---",
        "",
        "## Urutan Prioritas",
        "",
    ]

    for p in plans:
        if p.notes:
            lines.append(f"- **{p.bank} ...{p.last4}**: {p.notes}")

    lines += [
        "",
        "---",
        f"*Auto-generated {run_date}. Update checklist setelah setiap pembayaran.*",
    ]

    return "\n".join(lines)


def generate_payment_plan_md(
    plans: list[CardPaymentPlan],
    year_month: str,
    run_date: str,
    next_months: int = 3,
) -> str:
    label = _month_label(year_month)
    total_bal = sum(p.current_balance for p in plans)

    lines = [
        f"# Rencana Pembayaran Kartu Kredit — {label}",
        f"**SUWANDHY PRAHARTO** | Diperbarui: {run_date} | Total Hutang: **{_fmt(total_bal)}**",
        "",
        "---",
        "",
        f"## Saldo & Rencana Pembayaran {label}",
        "",
        "| Kartu | Saldo | Rekomendasi | JT | Estimasi Saldo Bulan Depan |",
        "|---|---|---|---|---|",
    ]

    for p in plans:
        est_next = max(0, p.current_balance + int(p.current_balance * 0.0175) - p.recommended_payment)
        est_tag = " ⚠️" if p.is_estimated else ""
        due = p.due_date or "—"
        lines.append(
            f"| {p.bank} ...{p.last4}{est_tag} | {_fmt(p.current_balance)} "
            f"| {_fmt(p.recommended_payment)} | {due} | ~{_fmt(est_next)} |"
        )

    lines += [
        "",
        "---",
        "",
        "## Cicilan Aktif Bulan Ini",
        "",
        "| Kartu | Deskripsi | Cicilan/Bln | Berakhir |",
        "|---|---|---|---|",
    ]

    for p in plans:
        insts = get_active_installments_for_card(p.card_id, year_month)
        for inst in insts:
            lines.append(
                f"| {p.bank} ...{p.last4} | {inst['description']} "
                f"| {_fmt(inst['monthly'])} | {inst['end_ym']} |"
            )

    lines += [
        "",
        "---",
        f"*Auto-generated {run_date}*",
    ]

    return "\n".join(lines)


def generate_excel(
    plans: list[CardPaymentPlan],
    year_month: str,
    output_path: str,
):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Pembayaran Bulanan"

    label = _month_label(year_month)
    budget = MONTHLY_BUDGET.get(year_month, DEFAULT_BUDGET)

    # Header
    headers = ["Bank", "Kartu", "Saldo (Rp)", "Min. Payment (Rp)",
               "Rekomendasi (Rp)", "Jatuh Tempo", "Prioritas", "Estimasi?", "Keterangan"]
    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(color="FFFFFF", bold=True)

    ws.append([f"Laporan {label} — SUWANDHY PRAHARTO"])
    ws.append([f"Total Hutang", sum(p.current_balance for p in plans),
               "Budget", budget])
    ws.append([])
    ws.append(headers)

    for cell in ws[4]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    priority_colors = {
        "KRITIS": "FF0000",
        "TINGGI": "FF6600",
        "SEDANG": "FFC000",
        "RENDAH": "92D050",
    }

    for p in plans:
        row = [
            p.bank,
            f"...{p.last4}",
            p.current_balance,
            p.minimum_payment,
            p.recommended_payment,
            p.due_date or "—",
            p.priority,
            "Ya" if p.is_estimated else "Tidak",
            p.notes,
        ]
        ws.append(row)
        color = priority_colors.get(p.priority, "FFFFFF")
        ws.cell(ws.max_row, 7).fill = PatternFill("solid", fgColor=color)

    # Total row
    ws.append([
        "TOTAL", "",
        sum(p.current_balance for p in plans),
        sum(p.minimum_payment for p in plans),
        sum(p.recommended_payment for p in plans),
        "", "", "", ""
    ])
    total_row = ws.max_row
    for col in range(1, 6):
        ws.cell(total_row, col).font = Font(bold=True)

    # Column widths
    ws.column_dimensions["A"].width = 15
    ws.column_dimensions["B"].width = 10
    for col in ["C", "D", "E"]:
        ws.column_dimensions[col].width = 18
    ws.column_dimensions["F"].width = 14
    ws.column_dimensions["G"].width = 10
    ws.column_dimensions["I"].width = 35

    wb.save(output_path)
